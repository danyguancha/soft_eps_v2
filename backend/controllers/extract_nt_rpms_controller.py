import pandas as pd
import os
import re
import math
import openpyxl
from datetime import datetime, date as date_type
from typing import List, Dict, Optional
from difflib import SequenceMatcher
from functools import lru_cache
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass


# ── Regex pre-compilados a nivel de módulo ────────────────────────────────────
_RE_WHITESPACE  = re.compile(r'\s+')
_RE_NORM_HEADER = re.compile(r'\s+')


# ========== CONFIGURACIÓN ==========
@dataclass
class SheetConfig:
    sheet_pattern:      str
    header_row_general: int
    header_row_data:    int
    data_start_row:     int
    output_label:       str


NT_RPMS_CONFIG = SheetConfig(r"NT\s*RPMS",    4, 11, 12, "NT_RPMS")
NT_RMPN_CONFIG = SheetConfig(r"NT[_\s]*RMPN", 8, 19, 20, "NT_RMPN")


# ========== CACHÉ GEOGRÁFICO ==========
_GEOGRAPHIC_CACHE: Dict = {}


@lru_cache(maxsize=1)
def load_geographic_data_cached(file_path: str) -> pd.DataFrame:
    if file_path in _GEOGRAPHIC_CACHE:
        return _GEOGRAPHIC_CACHE[file_path].copy()
    if not os.path.exists(file_path):
        print(f"Archivo geográfico no encontrado: {file_path}")
        return pd.DataFrame()
    try:
        print(f"📍 Cargando datos geográficos: {file_path}")
        df = pd.read_excel(file_path, engine='openpyxl', dtype=str, na_filter=False)
        df.columns = df.columns.str.strip()
        _GEOGRAPHIC_CACHE[file_path] = df
        print("✓ Datos geográficos cargados y almacenados en caché")
        return df.copy()
    except Exception as e:
        print(f"Error cargando datos geográficos: {e}")
        return pd.DataFrame()


# ========== UTILIDADES ==========
def find_target_sheet(sheet_names: List[str], pattern: str) -> Optional[str]:
    for name in sheet_names:
        if re.search(pattern, name, re.IGNORECASE):
            return name
    return sheet_names[0] if sheet_names else None


def _normalize_header(text: str) -> str:
    return _RE_NORM_HEADER.sub(' ', str(text).strip()).lower()


def similarity_score_strict(t1: str, t2: str) -> float:
    if not t1 or not t2: return 0.0
    t1, t2 = _normalize_header(t1), _normalize_header(t2)
    if t1 == t2: return 1.0
    return SequenceMatcher(None, t1, t2).ratio()


def _collapse_spaces(s: str) -> str:
    return _RE_WHITESPACE.sub(' ', s).strip()


# ========== DISPLAY VALUE DE CELDA ==========
def get_cell_display_value(cell) -> Optional[str]:
    """
    Compatible con Cell, ReadOnlyCell y EmptyCell de openpyxl.
    EmptyCell.value siempre es None → retorna None de inmediato.
    Usa getattr para number_format por si EmptyCell no lo expone.
    """
    value = cell.value
    if value is None:
        return None

    if isinstance(value, str):
        s = _collapse_spaces(value)
        return s if s and s.lower() != 'nan' else None

    if isinstance(value, bool):
        return 'VERDADERO' if value else 'FALSO'

    if isinstance(value, (datetime, date_type)):
        return value.strftime('%d/%m/%Y')

    if isinstance(value, (int, float)):
        if isinstance(value, float) and math.isnan(value):
            return None

        # CORRECCIÓN: getattr protege contra EmptyCell sin number_format
        fmt = getattr(cell, 'number_format', None) or 'General'

        if '%' in fmt:
            pct = value * 100
            if abs(pct - round(pct)) < 1e-6:
                return f"{int(round(pct))}%"
            return f"{pct:.2f}".replace('.', ',') + "%"

        if '$' in fmt:
            int_val = int(round(value))
            miles = f"{abs(int_val):,}".replace(',', '.')
            return f"${miles}" if int_val >= 0 else f"-${miles}"

        if isinstance(value, int) or (isinstance(value, float) and value == int(value)):
            return str(int(value))

        s = f"{value:.10f}".rstrip('0').rstrip('.')
        return s.replace('.', ',')

    s = _collapse_spaces(str(value))
    return s if s and s.lower() != 'nan' else None


# ========== MAPEO DE ENCABEZADOS ==========
class HeaderMapper:
    _PO_KEYS = ['POBLACION OBJETO', 'Poblacion Objeto', 'poblacion objeto', 'poblacion_objeto']

    KEYWORDS: Dict[str, List[str]] = {
        'departamento':                 ['DPTO', 'dpto', 'departamento', 'depto'],
        'municipio':                    ['MUNICIPIO', 'municipio', 'mpio'],
        'nombre_ips':                   ['NOMBRE _ IPS', 'NOMBRE_IPS', 'nombre ips', 'ips', 'institucion'],
        'regimen':                      ['REGIMEN', 'regimen', 'régimen'],
        'proyeccion_tiempo':            ['Proyeccion de Tiempo de NT', 'proyeccion tiempo de nt', 'tiempo de nt'],
        'consultas_procedimientos':     ['CONSULTAS / PROCEDIMIENTOS', 'CONSULTAS/PROCEDIMIENTOS', 'consultas procedimientos'],
        'servicios_habilitados':        ['S_HABILITADOS', 'servicios habilitados', 'tipo de interven'],

        'frecuencia_edad':              [
            'FRECUENCIA SEGÚN EDAD',
            'FRECUENCIA SEGUN EDAD',
            'FRECUENCIA SEGÚN EDAD (MESES Ó AÑOS)',
            'FRECUENCIA SEGUN EDAD (MESES O AÑOS)',
            'FRECUENCIA SEGUN EDAD (MESES O ANOS)',
            'frecuencia según edad',
            'frecuencia segun edad',
        ],

        'cups':                         ['CUPS(AP,AC, MED, OTROS_S)', 'CUPS(AP.AC. MED. OTROS_S)', 'cups', 'código cups', 'codigo cups'],

        'frecuencia_indicada':          [
            'FRECUENCIA INDICADA',
            'frecuencia indicada',
            'FREC INDICADA',
            'FREC. INDICADA',
        ],

        'periodo':                      ['PERIODO', 'periodo', 'período'],

        'frecuencia_uso':               [
            'FRECUENCIA DE USO_IPS',
            'FRECUENCIA  DE USO_IPS',
            'FRECUENCIA DE USO IPS',
            'FRECUENCIA  DE USO IPS',
            'frecuencia de uso_ips',
            'frecuencia de uso ips',
            'frecuencia de uso',
            'FREC DE USO',
            'FREC_USO',
        ],

        'frecuencia_ajustada':          [
            'FRECUENCIA AJUSTADA',
            'frecuencia ajustada',
            'FRECUENCIA_AJUSTADA',
            'FREC AJUSTADA',
            'FREC_AJUSTADA',
        ],

        'meta':                         ['%META', '%meta', 'meta', 'META'],

        'atenciones_realizar_anual':    [
            'ATENCIONES A REALIZAR ANUAL',
            'atenciones a realizar anual',
            'ATENCIONES A REALIZAR',
            'atenciones a realizar',
        ],

        'intervenciones_realizadas':    [
            'INTERVENCIONES REALIZADAS HISTORICO',
            'INTERVENCIONES REALIZADAS',
            'intervenciones realizadas historico',
            'intervenciones realizadas',
        ],
        'finalidad_cups_3374':          ['FINALIDAD_CUPS (3374)', 'Finalidad_CUPS (3374)', 'finalidad cups 3374'],
        'finalidad_cups_1036':          ['FINALIDAD_CUPS (1036)', 'Finalidad_CUPS (1036)', 'finalidad cups 1036'],
        'poblacion_objeto':             _PO_KEYS,
        'poblacion_objeto_2':           _PO_KEYS,
        'curso_de_vida':                ['MOMENTOS DEL CURSO DE VIDA', 'curso de vida', 'curso_vida'],
        'fase_atencion':                ['FASE DE ATENCIÓN', 'FASE DE ATENCION', 'fase de atencion'],
        'tipo_intervencion':            ['TIPO DE INTERVENCIÓN', 'TIPO DE INTERVENCION', 'tipo de intervencion'],
        'genero':                       ['GENERO', 'genero', 'género'],
        'cie10':                        ['CIE10', 'cie10', 'código cie10', 'codigo cie10'],
        'descripcion_actividad':        ['DESCRIPCION', 'descripcion de la actividad', 'descripcion_actividad'],
        'pob_susceptible_anual':        ['POBLACION SUSCEPTIBLE ANUAL', 'poblacion susceptible anual'],

        'atenciones_realizar_ajustada': [
            'ATENCIONES A REALIZAR AJUSTADA',
            'ATENCIONES A REALIZAR ANUAL AJUSTADA',
            'atenciones a realizar ajustada',
            'atenciones a realizar anual ajustada',
            'atenciones_realizar_anual_ajustada',
            'atenciones_realizar_ajustada',
            'ATENCIONES AJUSTADAS',
            'atenciones ajustadas',
        ],

        'pob_susceptible_mensual':      ['POBLACION SUSCEPTIBLE MENSUAL', 'ATENCIONES A REALIZAR MENSUAL'],
        'modalidad':                    ['MODALIDAD', 'modalidad', 'tipo de modalidad'],
        'tarifa':                       ['TARIFA', 'tarifa', 'valor tarifa'],
        'costo_por_ips':                ['COSTO TOTAL POR IPS', 'costo total por ips', 'costo por ips'],
    }

    # Keywords pre-normalizados una sola vez al cargar el módulo
    _KEYWORDS_NORM: Dict[str, List[str]] = {
        field: [_RE_NORM_HEADER.sub(' ', kw.strip()).lower() for kw in kws]
        for field, kws in KEYWORDS.items()
    }

    _FUZZY_THRESHOLD = 0.75

    @classmethod
    def find_all_columns(cls, headers: List, fields: List[str]) -> Dict[str, Optional[int]]:
        result: Dict[str, Optional[int]] = {f: None for f in fields}
        used: set = set()

        # Pre-normalizar headers una sola vez para esta llamada
        headers_norm = [
            _RE_NORM_HEADER.sub(' ', str(h).strip()).lower()
            if h is not None and str(h).strip() else None
            for h in headers
        ]

        # ── Pasada 1: exact match ─────────────────────────────────────────────
        for field in fields:
            kws_norm = cls._KEYWORDS_NORM.get(field, [])
            for idx, h_norm in enumerate(headers_norm):
                if idx in used or h_norm is None:
                    continue
                if h_norm in kws_norm:
                    result[field] = idx
                    used.add(idx)
                    break

        # ── Pasada 2: fuzzy para los no resueltos ─────────────────────────────
        for field in fields:
            if result[field] is not None:
                continue
            kws_norm = cls._KEYWORDS_NORM.get(field, [])
            best_idx, best_score = None, cls._FUZZY_THRESHOLD
            for idx, h_norm in enumerate(headers_norm):
                if idx in used or h_norm is None:
                    continue
                for kw_norm in kws_norm:
                    score = SequenceMatcher(None, kw_norm, h_norm).ratio()
                    if score > best_score:
                        best_score, best_idx = score, idx
            result[field] = best_idx
            if best_idx is not None:
                used.add(best_idx)

        return result


# ========== FORMATEO ==========
class DataFormatter:

    @staticmethod
    def format_codigo(value, n: int) -> str:
        if value is None: return ""
        if isinstance(value, float) and math.isnan(value): return ""
        s = str(value).strip().split('.')[0]
        return s.zfill(n) if s.isdigit() else s

    @classmethod
    def format_dataframe(cls, df: pd.DataFrame) -> pd.DataFrame:
        df = df.copy()
        if 'codigo_departamento' in df.columns:
            df['codigo_departamento'] = df['codigo_departamento'].apply(
                lambda x: cls.format_codigo(x, 2))
        if 'codigo_municipio' in df.columns:
            df['codigo_municipio'] = df['codigo_municipio'].apply(
                lambda x: cls.format_codigo(x, 3))
        print("   ✓ Valores pegados tal cual desde Excel (display values)")
        return df


# ========== ENRIQUECIMIENTO GEOGRÁFICO ==========
class GeographicEnricher:

    def __init__(self, departamentos_file: str):
        self.mapping = self._load_mapping(departamentos_file)

    def _load_mapping(self, file_path: str) -> pd.DataFrame:
        df = load_geographic_data_cached(file_path)
        if df.empty: return pd.DataFrame()
        try:
            col_map: Dict[str, str] = {}
            for col in df.columns:
                c = col.lower()
                if   'departamento' in c and 'residencia' in c:       col_map['dpto_nombre'] = col
                elif 'coddpto' in c or ('cod' in c and 'dpto' in c):  col_map['dpto_codigo'] = col
                elif 'municipio' in c and 'residencia' in c:          col_map['mpio_nombre'] = col
                elif 'codmpio' in c or ('cod' in c and 'mpio' in c):  col_map['mpio_codigo'] = col
            if len(col_map) < 4:
                print("Columnas insuficientes en archivo geográfico")
                return pd.DataFrame()
            mapping = pd.DataFrame({
                'codigo_departamento': df[col_map['dpto_codigo']].str.strip().str.zfill(2),
                'departamento':        df[col_map['dpto_nombre']].str.strip(),
                'codigo_municipio':    df[col_map['mpio_codigo']].str.strip().str.zfill(3),
                'municipio':           df[col_map['mpio_nombre']].str.strip()
            })
            mapping['key'] = mapping['codigo_departamento'] + '_' + mapping['codigo_municipio']
            mapping = mapping.drop_duplicates(subset=['key'])
            print(f"✓ Preparados {len(mapping)} registros geográficos únicos")
            return mapping
        except Exception as e:
            print(f"Error procesando datos geográficos: {e}")
            return pd.DataFrame()

    def enrich(self, df: pd.DataFrame) -> pd.DataFrame:
        if self.mapping.empty: return df
        df = df.copy()
        df['_mpio_norm'] = df.apply(
            lambda r: '001'
            if str(r.get('codigo_departamento', '')).strip().zfill(2) == '91'
               and str(r.get('codigo_municipio', '')).strip().zfill(3) == '000'
            else str(r.get('codigo_municipio', '')).strip().zfill(3),
            axis=1
        )
        df['key'] = (df['codigo_departamento'].astype(str).str.zfill(2) + '_' +
                     df['_mpio_norm'].astype(str).str.zfill(3))
        df['codigo_municipio'] = df['_mpio_norm']
        df = df.merge(self.mapping[['key', 'departamento', 'municipio']],
                      on='key', how='left', suffixes=('_codigo', '_nombre'))
        if 'departamento_nombre' in df.columns:
            df = df.rename(columns={
                'departamento_codigo': 'codigo_departamento',
                'departamento_nombre': 'departamento',
                'municipio_codigo':    'codigo_municipio',
                'municipio_nombre':    'municipio'
            })
        df = df.drop(columns=['key', '_mpio_norm'], errors='ignore')
        enriched = df['departamento'].notna().sum()
        print(f"✓ Enriquecidos: {enriched}/{len(df)} registros ({enriched/len(df)*100:.1f}%)")
        return df


# ========== EXTRACTOR DE HOJAS ==========
class SheetExtractor:
    """
    ⚡ OPTIMIZADO:
      - read_only=True + data_only=True: lectura en streaming, 3-5x más rápido
      - enumerate para calcular número de fila: nunca usa cell.row (EmptyCell no lo tiene)
      - getattr(cell, 'number_format', None): protege contra EmptyCell
      - Solo itera desde la primera fila necesaria
      - Terminación anticipada al detectar N filas vacías consecutivas
      - field_indices pre-calculado fuera del bucle de datos
    """
    _MAX_CONSECUTIVE_EMPTY = 15
    _GENERAL_FIELDS = ['departamento', 'municipio', 'nombre_ips', 'regimen', 'proyeccion_tiempo']

    _DATA_FIELDS = [
        'consultas_procedimientos', 'servicios_habilitados', 'frecuencia_edad', 'cups',
        'frecuencia_indicada', 'periodo', 'frecuencia_uso', 'frecuencia_ajustada',
        'meta', 'atenciones_realizar_anual', 'intervenciones_realizadas',
        'finalidad_cups_3374', 'finalidad_cups_1036',
        'poblacion_objeto', 'poblacion_objeto_2',
        'curso_de_vida', 'fase_atencion', 'tipo_intervencion',
        'genero', 'cie10', 'descripcion_actividad',
        'pob_susceptible_anual', 'atenciones_realizar_ajustada', 'pob_susceptible_mensual',
        'modalidad', 'tarifa', 'costo_por_ips',
    ]

    def __init__(self, file_path: str, config: SheetConfig):
        self.file_path = file_path
        self.config    = config

    def extract(self) -> pd.DataFrame:
        wb = openpyxl.load_workbook(self.file_path, data_only=True, read_only=True)
        sheet_name = find_target_sheet(wb.sheetnames, self.config.sheet_pattern)
        if not sheet_name:
            wb.close()
            raise ValueError(
                f"No se encontró hoja '{self.config.sheet_pattern}' en {self.file_path}")
        print(f"      📄 Hoja encontrada: '{sheet_name}'")
        ws = wb[sheet_name]

        gen_hdr_1b  = self.config.header_row_general
        gen_val_1b  = self.config.header_row_general + 1
        data_hdr_1b = self.config.header_row_data
        data_str_1b = self.config.data_start_row
        start_row   = min(gen_hdr_1b, data_hdr_1b)

        gen_headers_raw:  Optional[List] = None
        gen_values_raw:   Optional[List] = None
        data_headers_raw: Optional[List] = None
        data_rows_raw:    List[List]     = []
        empty_streak = 0

        # CORRECCIÓN PRINCIPAL: enumerate calcula rn sin tocar cell.row
        # EmptyCell no tiene .row → acceder a él lanzaba AttributeError
        for row_offset, row_cells in enumerate(ws.iter_rows(min_row=start_row)):
            rn = start_row + row_offset  # número de fila 1-based, calculado externamente

            if not row_cells:
                continue

            if rn == gen_hdr_1b:
                gen_headers_raw = [get_cell_display_value(c) for c in row_cells]

            elif rn == gen_val_1b:
                gen_values_raw = [get_cell_display_value(c) for c in row_cells]

            elif rn == data_hdr_1b:
                data_headers_raw = [get_cell_display_value(c) for c in row_cells]

            elif rn >= data_str_1b:
                vals = [get_cell_display_value(c) for c in row_cells]
                if any(v is not None for v in vals):
                    data_rows_raw.append(vals)
                    empty_streak = 0
                else:
                    empty_streak += 1
                    if empty_streak >= self._MAX_CONSECUTIVE_EMPTY:
                        break

        wb.close()

        if gen_headers_raw is None or gen_values_raw is None:
            raise ValueError("No se encontraron filas de encabezado general")
        if data_headers_raw is None:
            raise ValueError("No se encontró fila de encabezado de datos")
        if not data_rows_raw:
            raise ValueError(f"No se encontraron datos válidos en hoja '{sheet_name}'")

        general = self._extract_general_from_raw(gen_headers_raw, gen_values_raw)
        rows    = self._extract_data_rows_from_raw(data_headers_raw, data_rows_raw, general)

        if not rows:
            raise ValueError(f"No se encontraron filas con datos válidos en hoja '{sheet_name}'")
        return pd.DataFrame(rows)

    def _extract_general_from_raw(self, headers: List, values: List) -> Dict:
        cols    = HeaderMapper.find_all_columns(headers, self._GENERAL_FIELDS)
        max_idx = len(values)

        def get(key):
            idx = cols.get(key)
            return values[idx] if (idx is not None and idx < max_idx) else None

        return {
            'codigo_departamento': get('departamento'),
            'codigo_municipio':    get('municipio'),
            'nombre_ips':          get('nombre_ips'),
            'regimen':             get('regimen'),
            'proyeccion_tiempo':   get('proyeccion_tiempo'),
        }

    def _extract_data_rows_from_raw(
        self,
        headers: List,
        data_rows_raw: List[List],
        general: Dict
    ) -> List[Dict]:
        cols = HeaderMapper.find_all_columns(headers, self._DATA_FIELDS)

        print("      🗂  Mapeo de columnas detectado:")
        for field, idx in cols.items():
            label = (f"col {idx} → '{headers[idx]}'"
                     if idx is not None and idx < len(headers) else "NO ENCONTRADA")
            print(f"         {field:<35} {label}")

        # Pre-calcular índices fuera del bucle de filas
        field_indices = [(field, cols.get(field)) for field in self._DATA_FIELDS]

        rows = []
        for raw_row in data_rows_raw:
            row_len = len(raw_row)
            row = general.copy()
            for field, idx in field_indices:
                row[field] = raw_row[idx] if (idx is not None and idx < row_len) else None
            if row.get('consultas_procedimientos') is not None or row.get('cups') is not None:
                rows.append(row)
        return rows


# ========== PROCESADOR PRINCIPAL ==========
class ExcelProcessor:

    COLUMN_ORDER = [
        'nombre_archivo', 'codigo_departamento', 'departamento', 'codigo_municipio',
        'municipio', 'nombre_ips', 'regimen', 'proyeccion_tiempo',
        'consultas_procedimientos', 'servicios_habilitados', 'frecuencia_edad', 'cups',
        'frecuencia_indicada', 'periodo', 'frecuencia_uso', 'frecuencia_ajustada',
        'meta', 'atenciones_realizar_anual', 'intervenciones_realizadas',
        'finalidad_cups_3374', 'finalidad_cups_1036',
        'poblacion_objeto', 'poblacion_objeto_2',
        'curso_de_vida', 'fase_atencion', 'tipo_intervencion',
        'genero', 'cie10', 'descripcion_actividad',
        'pob_susceptible_anual', 'atenciones_realizar_ajustada', 'pob_susceptible_mensual',
        'modalidad', 'tarifa', 'costo_por_ips',
    ]

    def __init__(self, folder_path: str, config: SheetConfig,
                 departamentos_file: Optional[str] = None,
                 max_workers: Optional[int] = None):
        self.folder_path = folder_path
        self.config      = config
        self.results: List[pd.DataFrame] = []
        self.errors:  List[tuple]        = []
        cpu = os.cpu_count() or 2
        self.max_workers = max_workers or min(max(6, cpu), 16)
        self.enricher = (GeographicEnricher(departamentos_file)
                         if departamentos_file and os.path.exists(departamentos_file) else None)

    def _process_single_file(self, filename: str) -> Optional[pd.DataFrame]:
        try:
            df = SheetExtractor(
                os.path.join(self.folder_path, filename), self.config).extract()
            df.insert(0, 'nombre_archivo', filename)
            print(f"✓ {filename} - OK ({len(df)} registros)")
            return df
        except Exception as e:
            self.errors.append((filename, str(e)))
            print(f"✗ {filename} - Error: {e}")
            return None

    def process_folder(self) -> pd.DataFrame:
        if not os.path.isdir(self.folder_path):
            raise FileNotFoundError(f"Carpeta no existe: {self.folder_path}")
        excel_files = sorted(
            f for f in os.listdir(self.folder_path)
            if f.lower().endswith(('.xlsx', '.xls')) and not f.startswith('~$')
        )
        if not excel_files:
            raise ValueError(f"No hay archivos Excel en: {self.folder_path}")

        workers = min(self.max_workers, len(excel_files))
        print(f"🔄 [{self.config.output_label}] Procesando {len(excel_files)} archivos "
              f"con {workers} hilos...")

        with ThreadPoolExecutor(max_workers=workers) as executor:
            futures = {executor.submit(self._process_single_file, f): f for f in excel_files}
            for future in as_completed(futures):
                df = future.result()
                if df is not None:
                    self.results.append(df)

        if not self.results:
            raise ValueError(
                f"[{self.config.output_label}] No se procesó ningún archivo correctamente")

        all_cols = list(dict.fromkeys(col for df in self.results for col in df.columns))
        combined = pd.concat(
            [df.reindex(columns=all_cols) for df in self.results],
            ignore_index=True
        )

        if self.enricher:
            combined = self.enricher.enrich(combined)

        combined = DataFormatter.format_dataframe(combined)
        return self._order_columns(combined)

    def _order_columns(self, df: pd.DataFrame) -> pd.DataFrame:
        available = [c for c in self.COLUMN_ORDER if c in df.columns]
        return df[available + [c for c in df.columns if c not in self.COLUMN_ORDER]]

    def get_summary(self) -> Dict:
        return {
            "archivos_procesados":  len(self.results),
            "archivos_con_errores": len(self.errors),
            "total_registros":      sum(len(d) for d in self.results),
            "errores":              self.errors
        }


# ========== FUNCIÓN BASE ==========
def _extract_sheet_to_csv(
    folder_path: str,
    output_csv_path: str,
    config: SheetConfig,
    separator: str = ';',
    departamentos_file: Optional[str] = None
) -> Dict:
    processor = ExcelProcessor(folder_path, config, departamentos_file)
    df = processor.process_folder()
    output_dir = os.path.dirname(output_csv_path)
    if output_dir and not os.path.exists(output_dir):
        os.makedirs(output_dir)

    df.to_csv(
        output_csv_path,
        index=False,
        sep=separator,
        encoding='utf-8-sig',
        quoting=1
    )

    print(f"   ✓ [{config.output_label}] CSV generado: {output_csv_path}")
    print(f"{'='*60}\n")
    return {
        "success":                   True,
        "csv_path":                  output_csv_path,
        "total_rows":                len(df),
        "total_columns":             len(df.columns),
        "column_order":              df.columns.tolist(),
        "summary":                   processor.get_summary(),
        "has_geographic_enrichment": departamentos_file is not None,
        "sheet_label":               config.output_label
    }


# ========== FUNCIONES PÚBLICAS ==========
def extract_nt_rpms_to_csv(
    folder_path: str,
    output_csv_path: str,
    separator: str = ';',
    departamentos_file: Optional[str] = None,
    config: SheetConfig = NT_RPMS_CONFIG
) -> Dict:
    print(f"\n{'='*60}\nEXTRACCIÓN HOJA: {config.output_label}\n{'='*60}")
    return _extract_sheet_to_csv(
        folder_path, output_csv_path, config, separator, departamentos_file)


def extract_nt_rmpn_to_csv(
    folder_path: str,
    output_csv_path: str,
    separator: str = ';',
    departamentos_file: Optional[str] = None,
    config: SheetConfig = NT_RMPN_CONFIG
) -> Dict:
    print(f"\n{'='*60}\nEXTRACCIÓN HOJA: {config.output_label}\n{'='*60}")
    return _extract_sheet_to_csv(
        folder_path, output_csv_path, config, separator, departamentos_file)