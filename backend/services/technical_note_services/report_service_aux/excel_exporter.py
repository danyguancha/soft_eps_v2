# services/technical_note_services/report_service_aux/excel_exporter.py
from typing import Any, Dict, List
import io
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter


class ExcelExporter:
    """
    EXPORTADOR EXCEL con colores y headers combinados
    - Fila 1: Headers combinados (meses, trimestres, semestres, anual)
    - Fila 2: Sub-headers (Numerador, Denominador, Cobertura %, Calificación)
    - Celdas de Calificación coloreadas según estado
    """
    
    # Paleta de colores para períodos
    COLORS = {
        'basico': 'F0F0F0',           # Gris claro para columnas básicas
        'enero': 'E8F4F8',            # Azul claro (Q1)
        'febrero': 'E8F4F8',
        'marzo': 'E8F4F8',
        'trimestre1': 'B3E5FC',       # Azul más oscuro (T1)
        'abril': 'F0F8E8',            # Verde claro (Q2)
        'mayo': 'F0F8E8',
        'junio': 'F0F8E8',
        'trimestre2': 'C8E6C9',       # Verde más oscuro (T2)
        'semestre1': 'A5D6A7',        # Verde oscuro (S1)
        'julio': 'FFF8E8',            # Naranja claro (Q3)
        'agosto': 'FFF8E8',
        'septiembre': 'FFF8E8',
        'trimestre3': 'FFE0B2',       # Naranja más oscuro (T3)
        'octubre': 'F8E8F8',          # Magenta claro (Q4)
        'noviembre': 'F8E8F8',
        'diciembre': 'F8E8F8',
        'trimestre4': 'E1BEE7',       # Magenta más oscuro (T4)
        'semestre2': 'CE93D8',        # Magenta oscuro (S2)
        'anual': 'FFCCBC'             # Rojo-naranja (Anual)
    }
    
    # 🚦 Colores de semaforización
    SEMAFORO_COLORS = {
        'NA': '808080',                    # Gris
        'Óptimo': '28a745',                # Verde
        'Aceptable': 'ffc107',             # Amarillo Claro
        'Deficiente': 'fd7e14',            # Amarillo Oscuro/Naranja
        'Muy Deficiente': 'ef1e1e',        # Rojo
        'Error': '6c757d'                  # Gris oscuro
    }
    
    def __init__(self):
        """Inicializa el exportador Excel"""
        self.meses_nombres = [
            "Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio",
            "Julio", "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre"
        ]
    
    def export_report(self, report_data: Dict[str, Any]) -> io.BytesIO:
        """Exporta reporte a Excel con colores y merges"""
        try:
            items = report_data.get('items', [])
            
            print(f"Exportando Excel: {len(items)} items")
            
            if not items:
                print("Sin datos para exportar")
                return None
            
            # Crear workbook
            wb = Workbook()
            ws = wb.active
            ws.title = "Reporte"
            
            # Generar headers con info de merges
            header_row2, color_mapping, merge_ranges = self._get_headers_info()
            
            # Escribir headers combinados
            self._write_headers_and_merges(ws, header_row2, color_mapping, merge_ranges)
            
            # Escribir datos con colores de Calificación
            self._write_data(ws, items, color_mapping, header_row2)
            
            # Ajustar ancho de columnas
            self._adjust_column_widths(ws)
            
            # Congelar las dos filas de headers
            ws.freeze_panes = 'A3'
            
            # Guardar en memoria
            excel_buffer = io.BytesIO()
            wb.save(excel_buffer)
            excel_buffer.seek(0)
            
            print(f"Excel generado en memoria: {len(items)} filas")
            return excel_buffer
            
        except Exception as e:
            print(f"Error exportando Excel: {e}")
            import traceback
            traceback.print_exc()
            return None
    
    def _get_headers_info(self) -> tuple:
        """
        Genera información de headers para dos niveles con merges
        
        Returns:
            Tupla (header_row2, color_mapping, merge_ranges)
        """
        # Headers de fila 2 (sub-headers atómicos)
        header_row2 = [
            "Procedimiento/Consulta", "Rango Edad", "Cups", "Frecuencia indicada", "Población Objeto", "Periodo", "Frecuencia de uso_ips",
            "Frecuencia ajustada anual", "Meta", "Pobl. Susceptible anual", "Pobl. Susceptible mensual",  "Proyección Tiempo"
        ]
        color_mapping = [self.COLORS['basico']] * len(header_row2)
        merge_ranges = []  # (nombre, col_inicio, col_fin)
        
        col = len(header_row2) + 1  # Siguiente columna disponible
        
        def add_period(name, color, subnames):
            """Helper para agregar un período con sus sub-columnas"""
            nonlocal col
            merge_init = col
            for sn in subnames:
                header_row2.append(sn)
                color_mapping.append(color)
                col += 1
            merge_ranges.append((name, merge_init, col - 1))
        
        # ========== ENERO, FEBRERO, MARZO → TRIMESTRE 1 ==========
        for idx in range(0, 3):
            add_period(
                self.meses_nombres[idx],
                self.COLORS[self.meses_nombres[idx].lower()],
                ["Numerador", "Denominador"]
            )
        
        add_period("Trimestre 1", self.COLORS['trimestre1'], 
                   ["Numerador", "Denominador", "Cobertura %", "Calificación"])
        
        # ========== ABRIL, MAYO, JUNIO → TRIMESTRE 2 → SEMESTRE 1 ==========
        for idx in range(3, 6):
            add_period(
                self.meses_nombres[idx],
                self.COLORS[self.meses_nombres[idx].lower()],
                ["Numerador", "Denominador"]
            )
        
        add_period("Trimestre 2", self.COLORS['trimestre2'], 
                   ["Numerador", "Denominador", "Cobertura %", "Calificación"])
        
        add_period("Semestre 1", self.COLORS['semestre1'], 
                   ["Numerador", "Denominador", "Cobertura %", "Calificación"])
        
        # ========== JULIO, AGOSTO, SEPTIEMBRE → TRIMESTRE 3 ==========
        for idx in range(6, 9):
            add_period(
                self.meses_nombres[idx],
                self.COLORS[self.meses_nombres[idx].lower()],
                ["Numerador", "Denominador"]
            )
        
        add_period("Trimestre 3", self.COLORS['trimestre3'], 
                   ["Numerador", "Denominador", "Cobertura %", "Calificación"])
        
        # ========== OCTUBRE, NOVIEMBRE, DICIEMBRE → TRIMESTRE 4 → SEMESTRE 2 → ANUAL ==========
        for idx in range(9, 12):
            add_period(
                self.meses_nombres[idx],
                self.COLORS[self.meses_nombres[idx].lower()],
                ["Numerador", "Denominador"]
            )
        
        add_period("Trimestre 4", self.COLORS['trimestre4'], 
                   ["Numerador", "Denominador", "Cobertura %", "Calificación"])
        
        add_period("Semestre 2", self.COLORS['semestre2'], 
                   ["Numerador", "Denominador", "Cobertura %", "Calificación"])
        
        add_period("Consolidado Anual", self.COLORS['anual'], 
                   ["Numerador", "Denominador", "Cobertura %", "Calificación"])
        
        return header_row2, color_mapping, merge_ranges
    
    def _write_headers_and_merges(self, ws, header_row2, color_mapping, merge_ranges):
        """Escribe headers con merges y estilos"""
        header_font = Font(bold=True, size=11, color='000000')
        header_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        thin_border = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin')
        )
        
        # Columnas básicas (merge vertical fila 1 y 2)
        headers_base = [
            "Procedimiento/Consulta", "Rango Edad", "Cups", "Frecuencia indicada", "Población Objeto", "Periodo", "Frecuencia de uso_ips",
            "Frecuencia ajustada anual", "Meta", "Pobl. Susceptible anual", "Pobl. Susceptible mensual",  "Proyección Tiempo"
        ]
        
        for i, header_name in enumerate(headers_base, 1):
            ws.merge_cells(start_row=1, start_column=i, end_row=2, end_column=i)
            cell = ws.cell(row=1, column=i)
            cell.value = header_name
            cell.font = header_font
            cell.fill = PatternFill(start_color=self.COLORS['basico'], 
                                   end_color=self.COLORS['basico'], fill_type='solid')
            cell.alignment = header_alignment
            cell.border = thin_border
        
        # Períodos combinados (merge horizontal en fila 1)
        for name, col_start, col_end in merge_ranges:
            # Merge fila 1 horizontalmente
            ws.merge_cells(start_row=1, start_column=col_start, end_row=1, end_column=col_end)
            cell = ws.cell(row=1, column=col_start)
            cell.value = name
            period_color = color_mapping[col_start - 1]
            cell.font = header_font
            cell.fill = PatternFill(start_color=period_color, end_color=period_color, fill_type='solid')
            cell.alignment = header_alignment
            cell.border = thin_border
            
            # Sub-headers en fila 2
            for col in range(col_start, col_end + 1):
                cell2 = ws.cell(row=2, column=col)
                cell2.value = header_row2[col - 1]
                cell2.font = header_font
                cell2.fill = PatternFill(start_color=color_mapping[col - 1], 
                                        end_color=color_mapping[col - 1], fill_type='solid')
                cell2.alignment = header_alignment
                cell2.border = thin_border
        
        # Altura de filas de headers
        ws.row_dimensions[1].height = 25
        ws.row_dimensions[2].height = 25
    
    def _write_data(self, ws, items, color_mapping, header_row2):
        """Escribe datos con colores de Calificación"""
        thin_border = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin')
        )
        data_alignment = Alignment(horizontal='center', vertical='center')
        
        # Identificar columnas de Calificación
        semaforo_cols = [i + 1 for i, h in enumerate(header_row2) if h == 'Calificación']
        
        row_start = 3
        for row_idx, item in enumerate(items):
            row_data = self._build_data_row(item)
            
            for col_num, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx + row_start, column=col_num)
                cell.value = value
                cell.alignment = data_alignment
                cell.border = thin_border
                
                # 🚦 Aplicar color de Calificación si es columna de Calificación
                if col_num in semaforo_cols:
                    semaforo_color = self._get_semaforo_color(value)
                    if semaforo_color:
                        cell.fill = PatternFill(start_color=semaforo_color, 
                                               end_color=semaforo_color, fill_type='solid')
                        cell.font = Font(bold=True, color='FFFFFF')  # Texto blanco
    
    def _get_semaforo_color(self, estado: str) -> str:
        """
        Obtiene el color hexadecimal según el estado del Calificación
        
        Args:
            estado: Texto del estado (Óptimo, Aceptable, etc.)
            
        Returns:
            Color hexadecimal sin '#'
        """
        return self.SEMAFORO_COLORS.get(estado, None)
    
    def _build_data_row(self, item) -> list:
        """Construye una fila de datos"""
        row = []
        
        # ========== COLUMNAS BÁSICAS ==========
        row.append(self._clean(item.get('consulta_procedimiento', '')))
        row.append(self._clean(item.get('rango_edad', '')))
        row.append(self._clean(item.get('cups', '')))        
        row.append(self._format_decimal(item.get('frecuencia_indicada', 0)))
        row.append(str(item.get('poblacion_objeto', 0)))
        row.append(self._clean(item.get('periodo', '')))
        row.append(self._format_decimal(item.get('frecuencia_uso_ips', '')))
        row.append(self._format_decimal(item.get('fecuencia_ajustada_anual', '')))
        row.append(self._format_decimal(item.get('meta', 0)))
        row.append(str(item.get('poblacion_susceptible_anual', 0)))
        row.append(str(item.get('poblacion_susceptible_mensual', 0)))
        row.append(str(item.get('proyeccion_tiempo', 12)))
        
        
        # ========== ENERO, FEBRERO, MARZO ==========
        for i in range(0, 3):
            mes = self.meses_nombres[i].lower()
            mes_data = item.get(mes, {})
            row.append(str(mes_data.get('numerador', 0)))
            row.append(str(mes_data.get('denominador', 0)))
        
        # T1
        t1 = item.get('T1', {})
        row.extend([
            str(t1.get('numerador', 0)),
            str(t1.get('denominador', 0)),
            self._format_decimal(t1.get('cobertura', 0)),
            self._clean(t1.get('semaforizacion', ''))
        ])
        
        # ========== ABRIL, MAYO, JUNIO ==========
        for i in range(3, 6):
            mes = self.meses_nombres[i].lower()
            mes_data = item.get(mes, {})
            row.append(str(mes_data.get('numerador', 0)))
            row.append(str(mes_data.get('denominador', 0)))
        
        # T2
        t2 = item.get('T2', {})
        row.extend([
            str(t2.get('numerador', 0)),
            str(t2.get('denominador', 0)),
            self._format_decimal(t2.get('cobertura', 0)),
            self._clean(t2.get('semaforizacion', ''))
        ])
        
        # S1
        s1 = item.get('S1', {})
        row.extend([
            str(s1.get('numerador', 0)),
            str(s1.get('denominador', 0)),
            self._format_decimal(s1.get('cobertura', 0)),
            self._clean(s1.get('semaforizacion', ''))
        ])
        
        # ========== JULIO, AGOSTO, SEPTIEMBRE ==========
        for i in range(6, 9):
            mes = self.meses_nombres[i].lower()
            mes_data = item.get(mes, {})
            row.append(str(mes_data.get('numerador', 0)))
            row.append(str(mes_data.get('denominador', 0)))
        
        # T3
        t3 = item.get('T3', {})
        row.extend([
            str(t3.get('numerador', 0)),
            str(t3.get('denominador', 0)),
            self._format_decimal(t3.get('cobertura', 0)),
            self._clean(t3.get('semaforizacion', ''))
        ])
        
        # ========== OCTUBRE, NOVIEMBRE, DICIEMBRE ==========
        for i in range(9, 12):
            mes = self.meses_nombres[i].lower()
            mes_data = item.get(mes, {})
            row.append(str(mes_data.get('numerador', 0)))
            row.append(str(mes_data.get('denominador', 0)))
        
        # T4
        t4 = item.get('T4', {})
        row.extend([
            str(t4.get('numerador', 0)),
            str(t4.get('denominador', 0)),
            self._format_decimal(t4.get('cobertura', 0)),
            self._clean(t4.get('semaforizacion', ''))
        ])
        
        # S2
        s2 = item.get('S2', {})
        row.extend([
            str(s2.get('numerador', 0)),
            str(s2.get('denominador', 0)),
            self._format_decimal(s2.get('cobertura', 0)),
            self._clean(s2.get('semaforizacion', ''))
        ])
        
        # Anual
        anual = item.get('anual', {})
        row.extend([
            str(anual.get('numerador', 0)),
            str(anual.get('denominador', 0)),
            self._format_decimal(anual.get('cobertura', 0)),
            self._clean(anual.get('semaforizacion', ''))
        ])
        
        return row
    
    def _adjust_column_widths(self, ws):
        """Ajusta el ancho de las columnas automáticamente"""
        for column in ws.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)
            
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            
            adjusted_width = min(max_length + 2, 25)
            ws.column_dimensions[column_letter].width = adjusted_width
    
    def _clean(self, text: str) -> str:
        """Limpia texto"""
        if text is None:
            return ''
        return str(text).replace('\n', ' ').replace('\r', '')
    
    def _format_decimal(self, value: float) -> str:
        """Formatea decimal"""
        try:
            return str(round(float(value), 2))
        except (ValueError, TypeError):
            return '0'
